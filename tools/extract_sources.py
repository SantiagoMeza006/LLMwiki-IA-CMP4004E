"""Extract plain text from every file in raw/ into extracted/ so the LLM can grep sources.

Usage:  python tools/extract_sources.py            # only new/changed files
        python tools/extract_sources.py --all      # re-extract everything

Needs: python-pptx (pip install python-pptx) and pdftotext (poppler) on PATH.
raw/ is never modified. extracted/ is derived and can be deleted/regenerated at any time.
"""
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "raw")
OUT = os.path.join(ROOT, "extracted")


def pptx_to_text(path):
    import pptx
    lines = []
    for i, slide in enumerate(pptx.Presentation(path).slides, 1):
        lines.append(f"\n=== Slide {i} ===")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                lines.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[NOTES] {notes}")
    return "\n".join(lines)


def pdf_to_text(path):
    # Reading-order mode (no -layout) keeps two-column papers readable.
    return subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"],
                          capture_output=True, check=True).stdout.decode("utf-8", "replace")


def main():
    force = "--all" in sys.argv
    os.makedirs(OUT, exist_ok=True)
    for name in sorted(os.listdir(RAW)):
        src = os.path.join(RAW, name)
        base, ext = os.path.splitext(name)
        dst = os.path.join(OUT, base + ".txt")
        if not os.path.isfile(src) or ext.lower() not in (".pptx", ".pdf", ".md", ".txt"):
            continue
        if not force and os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src):
            continue
        if ext.lower() == ".pptx":
            text = pptx_to_text(src)
        elif ext.lower() == ".pdf":
            text = pdf_to_text(src)
        else:
            text = open(src, encoding="utf-8").read()
        with open(dst, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"extracted {name} -> extracted/{base}.txt ({len(text):,} chars)")


if __name__ == "__main__":
    main()
